from pathlib import Path
import json, statistics, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT.parent/'交付文件'
OUT.mkdir(exist_ok=True)
RES=ROOT/'experiments/results'
def load(name):return json.loads((RES/name).read_text(encoding='utf-8'))
def pct(x):return f'{100*x:.1f}%'
def f(x):return f'{x:.3f}'
retr={m:load('retrieval_'+m+'.json') for m in ['vector','hybrid','rerank']}
agent=load('agent_evaluation.json');summ=agent['summary']
chunk=load('chunking.json');emb=load('embeddings.json');parsers=load('parsers.json')
prompt=load('generation_parameters.json');robust=load('robustness.json')
manifest=json.loads((ROOT/'experiments/corpus_manifest.json').read_text(encoding='utf-8'))
baseline=load('agent_baseline_partial.json')['rows'][0]
final=next(r for r in agent['rows'] if r['id']==baseline['id'])
missing=sum(bool(r.get('metrics',{}).get('missing_citations')) for r in agent['rows'])
intro='本项目为南京农业大学课程实践题目一，研究方向为 Transformer 在计算机视觉中的应用。我们在原有 ragagent 仓库基础上迁移到本机 D 盘，补全索引、工具调用、实验和交付材料。所有数字来自保存的实测文件，不沿用任务说明中的示例提升比例。'
limits='独立人工评分仍待组员完成。引用页码存在仅说明来源能够定位，不等于该句受到原文支持。评测集包含初步人工编写与关键词辅助定位的答案和页码，正式评分前需复核标注。结果仅适用于当前论文集合和硬件。'
rt=[['检索模式','Hit@5','MRR@5','Recall@5','平均延迟 ms']]+[[m,pct(d['summary']['hit@5']),f(d['summary']['mrr@5']),pct(d['summary']['recall@5']),f(d['summary']['avg_latency_ms'])] for m,d in retr.items()]
ct=[['切分方式','块数','平均字符','Hit@5','MRR@5']]+[[d['summary']['strategy'],str(d['summary']['chunks']),f"{d['summary']['avg_length']:.1f}",pct(d['summary']['hit@5']),f(d['summary']['mrr@5'])] for d in chunk]
et=[['嵌入模型','维度','Hit@5','MRR@5','耗时 s']]+[[d['summary']['model'],str(d['summary']['dimensions']),pct(d['summary']['hit@5']),f(d['summary']['mrr@5']),f(d['summary']['elapsed_s'])] for d in emb]
docs=[]
def add(name,title,blocks):docs.append((name,title,blocks))
add('01_技术设计文档','智能科研助理论文知识库技术设计文档',[
 ('p',intro),('h','项目目标与验收范围'),('p','系统支持本地导入论文、带页码检索、多论文分析、计算工具和会话管理。设计重点是证据可追溯、模块可替换及错误可观察。联网搜索为可选能力，本版本未接入外部搜索服务。'),
 ('h','系统架构'),('p','Streamlit 接收文档和问题。加载器输出包含 doc_id、doc_name、物理页码和文本的记录，切分后经 bge-m3 转为向量并写入 Chroma。Agent 根据问题选择工具，检索工具返回证据，模型综合后以流式方式显示答案，轨迹和指标同步进入日志。'),
 ('table',[['层次','实现文件','责任'],['交互','ui/app.py','三列布局、会话、上传管理、轨迹和健康检查'],['决策','agent/react_loop.py','手写 ReAct、快速路由、并行与历史管理'],['工具','tools/registry.py','八个工具的签名、描述、执行函数'],['生成','rag/chain.py 与 rag/cache.py','上下文、引用、缓存、降级与日志'],['检索','retriever/api.py 与 store.py','范围过滤、增量索引、三种检索'],['基础','common/config.py 与 llm.py','统一配置、数据结构和本地模型客户端']]),
 ('h','数据加载与增量索引'),('p','PDF 默认用 PyMuPDF 提取并记录从 1 开始的物理页码，过滤参考文献后切分。DOCX、TXT、MD 作为逻辑页 1 处理。PDF 超过 300 页拒绝加载，网页限制单文件 32 MB。批量导入逐文件报告成功或失败，瞬态异常最多尝试三次，格式或路径错误直接返回。扫描文档没有 OCR，需要先获得可提取文本的 PDF。'),
 ('p','Chroma 按 chunk_id upsert，同一文档重导入会删除旧版本多余块。仅为新文本生成向量，SQLite 嵌入缓存按模型及文本哈希复用。索引记录嵌入模型和语料版本，拒绝把不同模型向量混在同一个库。Windows 原生 HNSW 在中文目录出现重启失败，因此采用 D:/NJAU_RAG_Data/index。'),
 ('h','向量数据库选型'),('table',[['维度','Chroma','FAISS'],['文档元数据与过滤','直接保存文档及页码，支持 where','需要额外映射和过滤逻辑'],['增量文档管理','upsert 和 delete 直接可用','需管理 ID 与删除策略'],['运行与性能','持久化适合课程规模','专注近邻搜索，适合大规模扩展'],['本项目决定','采用并实际运行','只做设计比较，未做性能对跑']]),
 ('h','检索与证据拼接'),('p','vector 使用归一化向量近邻。hybrid 将向量与 BM25 候选按 RRF 排名融合，当前权重 0.7 与 0.3。rerank 使用 bge-reranker-base 对候选重排。指定 doc_ids 在向量查询和 BM25 语料构建之前生效，多论文任务可分别检索后合并，避免小论文被候选集合挤掉。上下文按相关性排列并截断，证据标记采用【文档名-第X页】。'),
 ('table',rt),('p','本次实测 vector 的 Hit@5 最高，因此最终默认 vector，同时保留另外两种模式供复现。检索分数是模型相似度或重排分值，不是校准后的正确率。'),
 ('h','Agent 控制与工具'),('table',[['工具','用途'],['rag_search','返回知识库证据与来源'],['paper_meta','首页与已核验 manifest 的标题作者年份摘要 DOI'],['paper_compare','逐论文检索方法、数据集和结果'],['extract_keywords','提取关键词'],['summarize_paper','生成背景方法结果结论所需的证据'],['current_time','北京时间 UTC+08:00'],['calculator','受限 AST 运算'],['list_documents','查看真实 doc_id 和块数']]),
 ('p','同步和流式 API 共享事件引擎，最多四轮，每轮最多三个独立工具并行。计算、日期和列表采用确定路由，不消耗模型生成 Token。明确论文的元信息直接返回核验记录，其他明确论文任务先执行对应工具再生成；一般问题继续使用 JSON ReAct 规划。工具参数按函数签名校验，连接异常重试一次，超时跳过。线程超时不能强杀已运行的函数，后台调用仍可能短暂占用资源。'),
 ('h','缓存和会话记忆'),('p','语义缓存存入 SQLite，精确问题先查，再以向量余弦相似度 0.97 匹配。键包括语料版本、模型、检索模式、文档范围、会话和历史。数字及英文专业词变化禁止复用，TTL 为 24 小时，最多 1000 条。错误或漏引用的模型回答不写入 Agent 缓存。不同会话独立保存，历史预算按字符保守估算 3000，保留最近三轮并将较早内容摘要压缩。报告中的生成 Token 则来自 Ollama API 实际计数。'),
 ('h','引用控制与降级'),('p','错误页码从最终答案中替换为引用未核验。模型漏掉逐句引用时，附上真实检索来源并标明需要核对，missing_citations 指标为真且该回答不视为通过引用要求。无结果时可生成一般知识回答并标注未经知识库核验，低相关性时提示并展示候选，API 失败时给出重试建议。'),
 ('h','部署配置与当前边界'),('p','本机 Python 3.12、Ollama、Qwen2.5 3B、bge-m3、重排及对照模型均存放在 D 盘。16 GB 内存和 CPU 推理条件下采用 3B 以保证能运行，但不能期待大模型级别的理解质量。依赖通过 uv.lock 和 requirements-lock.txt 锁定。Dockerfile 和 Compose 提供容器方案，本机无 Docker，仅做配置静态核对。'),('p',limits)])
add('02_综合评测报告','智能科研助理综合评测报告',[
 ('p',intro),('h','数据与评测协议'),('p',f'语料共 {len(manifest)} 篇公开论文，合计 {sum(x["pages"] for x in manifest)} 个 PDF 物理页。检索集包含 68 道题，其中 36 道事实、12 道对比、6 道归纳、6 道推理、8 道库外问题。检索质量在 60 道可回答题上计算，8 道库外题另看拒答与 Top1 分布。Agent 选取其中 20 道代表题，加 4 道工具问题，实际运行 24 次。'),
 ('p','gold_sources 采用文档 ID 与物理页码组成的集合，不混淆不同论文同页。Hit@5 表示至少一个目标来源出现在前五块；MRR@5 是第一个目标来源排名的倒数；Recall@5 为目标来源召回比例；all_sources@5 要求多篇论文目标全部找回。相邻页含同样答案仍可能按严格页码算未命中。'),
 ('h','三种检索结果'),('table',rt),('p','重排序没有带来预期收益。当前候选、中文问题与英文证据、模型领域匹配和参考文献惩罚都可能影响结果，需要逐条分析，不能仅凭该实验确定单一因果。三种模式保留原始 Top5、问题、目标页及分数。'),
 ('p','延迟结果包含冷启动、嵌入缓存、重排加载以及同机其他实验竞争。rerank 第一题包含约 87 秒加载，因此平均延迟不是纯推理速度。只运行一次，不提供统计置信区间，也不以这些时间认定严格的硬件性能优劣。'),
 ('h','Agent 实测'),('table',[['指标','结果','口径'],['工具路由',pct(summ['tool_routing_accuracy']),'至少调用一个期望工具'],['平均迭代',f(summ['avg_iterations']),'包含确定路由，不表示模型推理轮次'],['平均端到端延迟',f(summ['avg_latency_ms']/1000)+' s','实际请求壁钟时间'],['平均 Token',f(summ['avg_tokens']),'Ollama 输入加输出计数'],['执行及引用门槛通过',pct(summ['success_rate']),'不能代替内容正确率'],['无效页码数',str(summ['invalid_citations']),'引用是否在观察来源中'],['漏逐句引用请求',str(missing)+' / 24','附来源提示，禁止缓存']]),
 ('p','两个对比任务未命中预期 paper_compare，而是分别调用 rag_search。仍可取得多文档资料，但按工具标签定义属于路由未命中。页码验证无法发现“引用存在但事实误读”的错误。独立人工质量评分尚未完成，不能宣称人工正确率或完整性达到某一比例。'),
 ('h','组件与生成参数'),('table',ct),('p','切分实验固定 Transformer 与 MAE 两篇、八道题，章节点切分 MRR 最高。样本不足以代表所有论文。切分长度单位为字符，不是 Token。'),('table',et),('p','三种嵌入使用相同 122 块及八道题，bge-large-zh-v1.5 查询使用检索指令前缀。bge-m3 命中嵌入缓存，其计时不可直接与两个本地模型的首次计算时间比较。'),
 ('table',[['温度','Top-p','Top-k','输出 Token','延迟 s']]+[[str(x['options']['temperature']),str(x['options']['top_p']),str(x['options']['top_k']),str(x['completion_tokens']),f(x['elapsed_ms']/1000)] for x in prompt]),('p','生成参数只比较一道 MAE 问题，保存四份真实回答供人工核对。它说明输出变化，不能建立总体质量排名。默认低温度 0.1，以减少随机变化。'),
 ('h','缓存与边缘验证'),('table',[['测试','结果']]+[[x['test'],('缓存命中 '+str(x['metrics']['cache_hit'])+'，'+str(x['metrics']['elapsed_ms'])+' ms') if 'cache_' in x['test'] else ('来源隔离 '+str(x.get('isolated'))) if x['test'].startswith('concurrent') else str(x.get('passed'))] for x in robust]),('p','缓存实测采用元信息问题及其中文改写，独立隔离验证采用 MAE 与 Swin 并发请求。缓存速度是这个具体样例的结果，不能声称所有问题都有同样加速。来源隔离只检查文档混入，不证明回答正确。27 项自动化测试覆盖并行、超时、签名、缓存失效、历史预算、多来源指标和引用检测。'),
 ('h','人工评审方法'),('p','使用 experiments/results/manual_review.csv，由未参与答案生成的组员对正确性、完整性、引用准确性各打 0 至 5 分，记录评审人和备注。正确性 0 为关键事实错误或无法回答，3 为主体正确但有局部偏差，5 为完全吻合原文。完整性检查问题要求是否全部覆盖。引用准确性必须打开论文对应页核对支持关系。建议两人独立评分，分差超过 1 分时复议。'),('p',limits),
 ('h','论文来源'),('table',[['文档','年份','arXiv ID']]+[[x['doc_id'],str(x['year']),x['arxiv_id']] for x in manifest]),('p','各论文的官方页面、PDF 链接、SHA256、页数及首页作者文本均保存在 experiments/corpus_manifest.json。上游 PDF 版本更新会改变页码和哈希，复现应优先使用交付包中的原始文件。')])
add('03_用户使用手册','智能科研助理安装与使用手册',[
 ('p','本手册说明当前电脑启动、日常论文管理、问答演示和另一台电脑复现。所有主要文件在 D 盘，应用地址为 http://127.0.0.1:8501。'),
 ('h','当前电脑快速启动'),('p','打开 D:/生产实习大作业/ragagent，双击 start_app.cmd。脚本检查本地 Ollama 与 Streamlit 服务，未启动时在后台启动。首次模型加载会比后续慢。出现报错先查看 data/logs/streamlit.err.log 与 ollama.err.log。不要移动或删除同级 runtime、models 及索引目录。'),
 ('image',ROOT/'docs/assets/ui_final.png'),('p','图示为本机实际计算工具演示，界面左侧管理资料，中间对话，右侧查看决策轨迹和指标。'),
 ('h','导入与管理文档'),('p','左侧选择 PDF、DOCX、TXT、MD 文件并上传，点击导入后查看逐文件进度。PDF 最大 300 页、单文件最大 32 MB。扫描 PDF 需要先 OCR。同名文档再次导入更新原有块，不用重建全部论文。删除文档会移除索引并使相关缓存失效。删除前确认已保留原始论文。'),
 ('h','提问与核验'),('p','可以选定文档范围，也可以在问题中写真实论文 ID，例如“MAE 的作者和年份是什么”“比较 ViT 与 DeiT 的训练方法区别”“计算 3.14*2.56”。引用使用 PDF 物理页码。读答案后应打开对应原文检查是否支持结论，尤其是数值、方法比较与因果解释。看到漏引用或引用未核验提示时，应缩小问题、重新提问或直接阅读原文。'),
 ('h','会话和指标'),('p','新建会话使历史独立，切换会话可继续之前的问题。右侧显示工具输入、观察、耗时和 Token。这里的 RAG 命中仅指工具返回带页码的证据，并不是正确率。真实检索准确率查综合评测报告。清空语义缓存会删除可复用回答，但不删除论文。健康检查验证模型服务、所需模型和向量库状态。'),
 ('h','另一台 Windows 电脑安装'),('p','将完整代码包解压到 D 盘，在 PowerShell 运行 powershell -ExecutionPolicy Bypass -File scripts/setup.ps1。脚本从官方源下载工具及模型，需要网络和足够磁盘空间。之后检查 .env 的 INDEX_DIR 为 ASCII 路径，使用 .venv/Scripts/python.exe scripts/ingest_corpus.py 导入论文，再双击 start_app.cmd。环境默认 Python 3.12，依赖由 uv.lock 锁定。若不包含论文，先运行 scripts/download_papers.py。'),
 ('h','检索和模型切换'),('p','在 .env 设置 RETRIEVAL_MODE=vector、hybrid 或 rerank，重启应用生效。当前默认 vector。重排模型首次从 Hugging Face 下载，可在 RERANKER_MODEL 指向本地目录。更换 EMBEDDING_MODEL 时必须使用新的 INDEX_DIR 并重新导入。生成模型在 OLLAMA_MODEL 配置，不要只改网页显示。'),
 ('h','Docker 部署'),('p','安装 Docker 后执行 docker compose up --build，模型初始化完成后执行 docker compose exec app python scripts/ingest_corpus.py。数据挂载在 ./data，模型使用命名卷，浏览器访问 localhost:8501。本机没有 Docker，此方案尚未实际构建验证。'),
 ('h','常见问题'),('table',[['现象','处理方法'],['页面打不开','运行启动脚本，检查 8501 端口和 Streamlit 日志'],['模型请求超时','等待当前 CPU 任务完成，健康检查后重试'],['向量库重启报错','Windows 改用 ASCII 路径，新建库并导入'],['某文档检索不到','确认真实 doc_id、导入成功及检索范围'],['问题超出知识库','按提示核验外部资料，不将一般知识当论文结论'],['引用在但内容错','查对应 PDF 页，人工核对后修订答案']]),
 ('h','演示顺序'),('p','建议依次展示健康检查、十二篇文档、元信息、论文问答、两论文对比、计算器、会话切换、缓存重复问题，最后用一次真实错误案例说明核验边界。完整实验命令见 README。')])
add('04_分块策略实验报告','论文分块策略对照实验报告',[
 ('p','本实验比较固定长度、递归、语义边界及章节切分。实验固定两篇论文和八道题，不修改运行中的 Chroma 索引。'),('h','控制条件'),('p','论文为 Transformer 与 MAE，题号 F01、F02、F03、F22、F23、F24、R01、R04。所有方案使用 bge-m3、归一化余弦检索、指定论文范围和 Top5。固定长度采用 256/50、512/100、1024/200 的字符窗口与重叠，其余方案窗口上限 1024、重叠 200。章节切分保留标题结构，语义切分按局部句段边界组织，不宣称实现训练式语义分割模型。'),
 ('h','实测结果'),('table',ct),('p','fixed256 和递归切分可达到 75% Hit@5，但会产生更多块或更低的首个命中排名。章节切分同为 75%，MRR@5 为 0.656，且块数 122，较固定 256 的 391 块更少。本项目继续采用章节策略，避免以小样本直接断言普遍最优。'),
 ('h','切分长度与计算量'),('table',[['策略','切分 s','嵌入 s','最短块字符']]+[[x['summary']['strategy'],f(x['summary']['split_s']),f(x['summary']['embedding_s']),str(x['summary']['min_length'])] for x in chunk]),('p','实验按顺序运行且存在嵌入缓存，section 的嵌入时间主要是缓存读取。recursive 第一次初始化分词相关依赖耗时计入切分时间，因此这组时间不是公平的冷启动速度排名。最短块仍存在短标题碎片，可作为后续合并短块的优化点。'),
 ('h','三种 PDF 解析器'),('table',[['解析器','页数','提取字符','耗时 s']]+[[x['parser'],str(x['pages']),str(x['characters']),f(x['elapsed_s'])] for x in parsers]),('p','PyMuPDF 在这两份可复制文本 PDF 上提取最快，因此作为默认加载器。字符数不同不代表 OCR 准确率或阅读顺序质量。未针对扫描文档、复杂表格或双栏顺序做独立准确性评分。'),('h','复现'),('p','运行 .venv/Scripts/python.exe experiments/component_exp.py。原始数据为 experiments/results/chunking.json 与 parsers.json，保留逐题 Top5 文档页码。')])
add('05_嵌入模型与生成参数报告','嵌入模型与生成参数对照报告',[
 ('p','本报告补充题目一要求的嵌入模型选型与生成参数对比。模型均已在本机实际运行，不用下载成功替代实验结果。'),('h','嵌入实验协议'),('p','在同样 Transformer 与 MAE 的 122 个章节块、同样八道题上，对比 bge-m3:567m、m3e-base 和 bge-large-zh-v1.5。所有向量归一化，使用点积进行余弦排序。bge-large-zh 查询按模型检索习惯加入中文检索指令，m3e 无前缀。'),('table',et),('p','bge-m3 和 bge-large-zh 都命中六题，bge-m3 的 MRR 较高。m3e 命中四题。基于当前中英文混合问题及英文论文，选择 bge-m3 作为默认模型。样本只有八道，结论为当前样例的工程选择，不能推广为模型能力榜单。'),('p','m3e-base 为 768 维，另外两个为 1024 维。不同模型即使维数一样也不共享语义空间。比较耗时受到 bge-m3 缓存及本地模型加载影响，不作公平速度结论。'),
 ('h','生成 Prompt 设计'),('p','Prompt 指明科研助理角色、用户问题、排序后证据和引用规则。证据是数据而非指令。要求证据不足时说明未知，论文事实段落保留【文件名-第X页】。Agent 最终综合使用独立 system 指令，错误页码校验，漏引用显示核验提醒。上下文预算和输出上限分别为 4096 与 512 Token 配置。'),
 ('h','参数实验'),('table',[['温度','Top-p','Top-k','输入 Token','输出 Token']]+[[str(x['options']['temperature']),str(x['options']['top_p']),str(x['options']['top_k']),str(x['prompt_tokens']),str(x['completion_tokens'])] for x in prompt]),('p','固定问题为 MAE 的 asymmetric encoder-decoder 如何工作，检索证据固定。四个配置均保留原始答案、响应时间和引用校验。在这次运行中输出长短有变化，但没有独立人工评分，不能认定某个温度准确率最高。默认 0.1 旨在降低随机波动。'),
 ('h','来源与复现'),('p','嵌入模型来自 https://huggingface.co/moka-ai/m3e-base 和 https://huggingface.co/BAAI/bge-large-zh-v1.5；重排序来自 https://huggingface.co/BAAI/bge-reranker-base；bge-m3 通过官方 Ollama 模型库安装。执行 experiments/embedding_exp.py 与 experiments/prompt_exp.py，查看 embeddings.json 和 generation_parameters.json。')])
add('06_BadCase与优化报告','论文问答错误案例与优化报告',[
 ('p','错误案例用于解释失败机制和改进效果，不把“返回了内容”视为“内容正确”。本报告分为检索、生成、决策、工程四类。'),
 ('h','检索错误'),('p','全库混合检索在多论文问题中容易只返回一篇论文的片段。已将 doc_ids 范围过滤移到检索前，并为多文档归纳独立检索。旧全局 hybrid 的 Hit@5 为 55.0%，指定范围的等权 hybrid 为 63.3%，提升 8.3 个百分点。这是范围策略的对照，不代表在用户没有提供文档范围时能获得同样提升。'),
 ('p','RRF 权重调到 0.7/0.3 后 Hit@5 仍为 63.3%，MRR 从 0.385 降到 0.374，未证实收益。重排序 Hit@5 为 56.7%，低于向量 71.7%。因此保留实验结果并选择 vector，不将任务说明的 85% 示例写作本组结果。'),
 ('h','生成错误'),('p','原始界面 MAE 问答把原文 narrower and shallower 表述为“更窄更深”，并省略页码。并发实测又出现把“可见 25% patches”误写成“4 个 patches”。这些错误说明来源在场仍可能发生语义误读，不能靠字符串页码校验认定事实正确。原始界面截图保存于 docs/assets/ui_baseline_badcase.png，实际回答保存在鲁棒性结果文件中。'),
 ('p',f'修复包括独立 system 引用指令、原始答案页码与观察来源的集合校验、漏引用告警及禁止缓存。在最终 24 次实测中仍有 {missing} 次漏逐句引用，说明规则没有消除模型质量问题。附上真实来源只是方便核验，不给原答案背书。未来应增加逐句支持判断、英文原句对照和更强模型对照。'),
 ('h','Agent 决策错误'),('p','优化前 F01 为重复检索及摘要调用，四轮迭代且 Token 较多。为明确论文问题增加工具路由，一次执行后直接综合；保留一般问题的 ReAct 规划。多论文独立任务最多三路并发，参数按签名校验。'),('table',[['同一问题 F01','优化前','最终版本'],['迭代',str(baseline['metrics']['iterations']),str(final['metrics']['iterations'])],['Token',str(baseline['metrics']['total_tokens']),str(final['metrics']['total_tokens'])],['端到端延迟 s',f(baseline['metrics']['elapsed_ms']/1000),f(final['metrics']['elapsed_ms']/1000)]]),('p','这是一对单样例对照，前后系统配置和硬件负载也变化，不将延迟差全部归因于路由。迭代和 Token 减少可以由日志直接核查。另有两个对比问题走分论文 rag_search 而非预期 paper_compare，仍保留为路由失败案例。'),
 ('h','工程故障'),('table',[['故障','修复','验证'],['localhost 模型请求经过代理返回 502','本地 Ollama Client trust_env=False','真实生成请求通过'],['中文索引目录重启失败','Chroma 移到 D 盘 ASCII 目录并重建','进程重开可读十二篇文档'],['虚假引用数量总为零','在过滤前解析原始引用再求差集','回归测试检测第 99 页'],['同页引用碰撞','文档名与页码双键','不同论文同页可区分'],['模型漏引用缓存风险','不缓存未满足引用条件的回答','最终缓存与隔离测试通过']]),
 ('h','后续改进优先级'),('p','优先进行独立人工复核并修正 gold 标注；随后对跨论文归纳按文档均衡分配证据和上下文；在相同硬件、无其他任务、冷暖状态一致的条件下重做延迟测试；最后比较更强生成模型及训练领域更接近的 reranker。当前错误已经记录，不宣称全部解决。')])
add('07_需求验收与提交清单','题目一需求验收与提交清单',[
 ('p','本清单将课程要求与代码、实测及交付文件逐项对应。待人工或环境条件完成的项目明确标注，便于提交前检查。'),
 ('table',[['要求','证据','状态'],['多格式加载与解析器对照','loader.py、parsers.json、robustness.json','已实现并测试'],['三种固定长度及递归语义章节','chunker.py、chunking.json、报告04','已实测，限两篇八题'],['bge-large-zh 与 m3e 对照','embeddings.json、报告05','已实测并含 bge-m3'],['Chroma/FAISS 选型及增量','store.py、技术设计文档','Chroma 实跑，FAISS 设计比较'],['三档检索评测','三个 retrieval_*.json','60 道可回答题实测'],['Prompt 与参数对比','generation_parameters.json、报告05','一道固定题四参数配置'],['流式引用缓存降级日志','rag、agent、data/logs','已实现，质量限制见报告'],['手写 ReAct 与八工具','agent/react_loop.py、tools/registry.py','已实现'],['并行恢复隔离历史','tests.xml、robustness.json','自动测试通过'],['三列 UI 与健康检查','ui/app.py、界面截图','本机运行与浏览器实测'],['10 至20论文与50+问答','12 PDF、manifest、68题eval_set','已完成初步标注'],['Agent 路由迭代延迟Token','agent_evaluation.json','24 道真实模型请求'],['正确完整引用人工评分','manual_review.csv','待独立组员填写'],['四类BadCase及优化','报告06与前后原始数据','已整理，未夸大收益'],['完整代码与Git历史','代码包、原历史、新本地提交','本地保留，未推送'],['技术设计评测报告与PPT','报告01至07和答辩PPT','本地文件交付'],['Docker与使用手册','Dockerfile、Compose、报告03','配置提供，容器未实跑']]),
 ('h','提交前由组员完成'),('p','填写封面学校、学院、专业班级、姓名、学号和指导教师。请使用自己的实际身份和分工，不填虚构人员。独立完成 manual_review.csv 三个质量评分维度和评审人，再将均值补入报告。按教师实际要求检查是否需要额外演示录像、平台证明、组内贡献说明或指定排版。'),
 ('h','材料定位'),('p','D:/生产实习大作业/交付文件 包含七份 Word 报告、答辩 PPT、完整提交 ZIP。D:/生产实习大作业/ragagent 是可继续修改的工作仓库。models 与 runtime 提供本机运行环境，模型不放入提交 ZIP；接收者按照 README 下载。'),
 ('h','复现顺序'),('p','先按使用手册安装和导入论文，再运行自动测试、三档检索、切分、嵌入、生成参数、Agent 和鲁棒性实验。为公平比较速度，顺序执行并记录冷暖状态。实验会更新结果文件，提交前备份本次原始数据。'),('p',limits)])

architecture_at=next(i for i,b in enumerate(docs[0][2]) if b==('h','数据加载与增量索引'))
docs[0][2][architecture_at:architecture_at]=[('image',ROOT/'docs/assets/architecture.png'),('p','架构图的原生可编辑版本位于答辩 PPT 第三页，显示离线文档处理和在线问答两条链路。')]
api_at=next(i for i,b in enumerate(docs[0][2]) if b==('h','部署配置与当前边界'))
docs[0][2][api_at:api_at]=[
 ('h','接口与数据结构'),
 ('table',[['函数','主要参数','返回结果'],['load_any','path 文件路径','含 doc_id 页码和文本的页面列表'],['ingest_with_retry','path attempts chunk_mode','成功块数或异常'],['retrieve_best','query topk doc_ids','RetrievedChunk 列表'],['rag_answer','query topk doc_ids options','answer citations metrics debug'],['react_loop','question history doc_ids session_id','answer trace metrics tool_stats success'],['react_loop_stream','同同步接口','thought action observation final_chunk done 事件'],['remove_doc','doc_id','删除文档并使索引版本变化']]),
 ('p','RetrievedChunk 字段为 chunk_id、doc_id、doc_name、page、text、score。AgentStep 包含 step_idx、thought、action、action_input、observation、elapsed_ms。流式 done 事件保存最终经过引用校验的答案，前端以它更新会话，不能只拼接未核验的生成片段。'),
 ('h','关键配置参数'),
 ('table',[['配置','本机值','作用'],['OLLAMA_MODEL','qwen2.5:3b','生成与规划模型'],['EMBEDDING_MODEL','bge-m3:567m','嵌入及检索'],['INDEX_DIR','D:/NJAU_RAG_Data/index','Windows ASCII 索引路径'],['RETRIEVAL_MODE','vector','vector hybrid rerank 可切换'],['CHUNK_MODE','section','章节切分'],['MODEL_CONTEXT','4096','模型上下文 Token 上限'],['MAX_OUTPUT_TOKENS','512','单次输出上限'],['CACHE_THRESHOLD','0.97','语义缓存阈值'],['LLM_TIMEOUT 与 TOOL_TIMEOUT','300 秒','客户端与工具等待上限']])]
insert_at=next(i for i,b in enumerate(docs[1][2]) if b==('h','Agent 实测'))
docs[1][2][insert_at:insert_at]=[
 ('h','Top1 分数分布自评'),
 ('image',ROOT/'docs/assets/top1_distribution.png'),
 ('p','可回答题与库外题的 Top1 分数分布存在重叠，单一阈值不能保证可靠拒答。向量相似度与重排 sigmoid 分数语义不同，不能横向当作概率比较。分布原始统计在 top1_distribution.json。样本规模 60 与 8 不等，图中为题目数量而非比例。')]
def set_font(style,size,bold=False):
 style.font.name='Microsoft YaHei';style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor(0,0,0)
 style.element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
def table(doc,rows):
 t=doc.add_table(rows=1, cols=len(rows[0]));t.autofit=False
 for i,row in enumerate(rows):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for cell,val in zip(cells,row):
   cell.text=str(val)
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.1
    for r in p.runs:r.font.size=Pt(9.5);r.bold=i==0
   tcPr=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');mar.append(e)
   tcPr.append(mar)
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');tcPr.append(shade)
  trPr=t.rows[i]._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');trPr.append(keep)
  if i==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'C8C8C8');borders.append(e)
 t._tbl.tblPr.append(borders)
 if rows[0]==['文档','年份','arXiv ID']:
  for row in t.rows:
   for cell in row.cells:
    for par in cell.paragraphs:
     par.paragraph_format.space_before=Pt(0);par.paragraph_format.space_after=Pt(0);par.paragraph_format.line_spacing=1
     for r in par.runs:r.font.size=Pt(9)
    for el in cell._tc.xpath('.//w:tcMar/w:top | .//w:tcMar/w:bottom'):el.set(qn('w:w'),'40')
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
for name,title,blocks in docs:
 doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.top_margin=sec.bottom_margin=Inches(.75);sec.left_margin=sec.right_margin=Inches(.8)
 for sty,size,bold in [('Normal',11,False),('Title',22,True),('Heading 1',15,True),('Heading 2',12,True)]:set_font(doc.styles[sty],size,bold)
 for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 doc.styles['Normal'].paragraph_format.line_spacing=1.15;doc.styles['Normal'].paragraph_format.space_after=Pt(5)
 doc.styles['Heading 1'].paragraph_format.space_before=Pt(13);doc.styles['Heading 1'].paragraph_format.space_after=Pt(7)
 doc.add_paragraph(title,'Title')
 doc.add_paragraph('南京农业大学 课程实践 题目一\n版本 2026年10月10日 本机实测交付版')
 doc.add_paragraph('组员姓名：____________    学号：____________\n专业班级：____________    指导教师：____________')
 for kind,value in blocks:
  if kind=='p':doc.add_paragraph(value)
  elif kind=='h':doc.add_paragraph(value,'Heading 1')
  elif kind=='table':table(doc,value)
  elif kind=='image' and value.exists():
   from PIL import Image
   iw,ih=Image.open(value).size
   doc.add_picture(str(value),width=Inches(min(6.7,3.8*iw/ih)))
 footer=sec.footer.paragraphs[0];footer.alignment=2
 run=footer.add_run('题目一  智能科研助理   ');run.font.size=Pt(8)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 doc.save(OUT/(name+'.docx'))
 md=['# '+title,'',intro if name.startswith('01') else '南京农业大学课程实践 题目一','']
 for kind,val in blocks:
  if kind=='h':md+=['## '+val,'']
  elif kind=='p':md+=[val,'']
  elif kind=='table':md+=[' | '.join(val[0]),' | '.join(['---']*len(val[0]))]+[' | '.join(map(str,row)) for row in val[1:]]+['']
  elif kind=='image':md+=[f'![本机实测](assets/{val.name})','']
 (ROOT/'docs'/(name+'.md')).write_text('\n'.join(md),encoding='utf-8')
 print(name,flush=True)
print('7 editable reports written')
